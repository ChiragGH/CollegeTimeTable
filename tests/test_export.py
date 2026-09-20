"""
Tests for the timetable export package.

Covers:
    - Time grid helpers (periods, lunch, block positions)
    - Conflict audit (C1 teacher, C2 room, C3 section/group, C6 parallel)
    - View builders: class, section, group, teacher, room
    - G1/G2 parallel pairing and clear display
    - CSV exporters (grid, flat, conflicts, unscheduled)
    - XLSX exporters (merged blocks, conflict sheet, styling structure)
    - PDF exporters (valid documents, page counts)
    - Export service (bundle, entity filter, filenames)
    - JSON → Placement rebuild (web restart fallback)
"""

import csv
from pathlib import Path

import pytest

from app.models.enums import ActivityType, Day, RoomType
from app.models.assignment import Assignment
from app.models.placement import Placement
from app.models.slot import TimeSlot
from app.models.teacher import Teacher
from app.models.room import Room
from app.models.subject import Subject
from app.models.section import Section

from app.export import (
    ExportContext,
    TimetableExporter,
    audit_placements,
    build_views,
    placements_from_result_data,
)
from app.export.grid import (
    COL_LUNCH,
    GRID_COLUMNS,
    PERIOD_TIMES,
    lunch_label,
    slots_are_consecutive,
)
from app.export.views import (
    VIEW_CLASS,
    VIEW_GROUP,
    VIEW_KINDS,
    VIEW_ROOM,
    VIEW_SECTION,
    VIEW_TEACHER,
)


# ====================================================================
# Fixtures & factories
# ====================================================================

def slot(day: str, period: int) -> TimeSlot:
    return TimeSlot(day=Day[day], period=period)


@pytest.fixture
def ctx() -> ExportContext:
    teachers = {
        "T001": Teacher("T001", "RJS", "CSE", True),
        "T002": Teacher("T002", "SRP", "CSE", True),
        "T003": Teacher("T003", "AMK", "CSE", True),
    }
    rooms = {
        "R001": Room("R001", "CC1", RoomType.LAB, "CSE", False, True),
        "R002": Room("R002", "CC2", RoomType.LAB, "CSE", False, True),
        "R010": Room("R010", "L5", RoomType.LECTURE, None, True, True),
        "R011": Room("R011", "L6", RoomType.LECTURE, None, True, True),
    }
    subjects = {
        "SUB1": Subject("SUB1", "3.2", "Data Structures", "CSE", 3, True),
        "SUB2": Subject("SUB2", "3.4", "Database Management Systems", "CSE", 3, True),
        "SUB3": Subject("SUB3", "3.7", "Data Structures Lab", "CSE", 3, True),
    }
    return ExportContext(
        session_id="2026-27_CSE_Sem3",
        academic_year="2026-27",
        branch="CSE",
        semester=3,
        sections=[Section("CSE", 3, "A"), Section("CSE", 3, "B")],
        teachers=teachers,
        rooms=rooms,
        subjects=subjects,
        college_name="Test Polytechnic",
        generated_at="01 Jan 2026, 00:00",
    )


def make_placement(
    placement_id: str,
    teacher_id: str,
    subject_id: str,
    section: str,
    group: str,
    activity: ActivityType,
    slots,
    room_id: str,
) -> Placement:
    slots = list(slots)
    assignment = Assignment(
        assignment_id=f"A_{placement_id}",
        session_id="test",
        teacher_id=teacher_id,
        subject_id=subject_id,
        branch="CSE",
        semester=3,
        section=section,
        group=group,
        activity_type=activity,
        weekly_periods=len(slots),
        room_id=room_id,
        block_size=len(slots),
        sessions_per_week=1,
    )
    return Placement(
        placement_id=placement_id,
        assignment=assignment,
        slots=slots,
        room_id=room_id,
    )


@pytest.fixture
def clean_placements() -> list:
    """A valid timetable snippet: lecture, lecture block, parallel G1/G2."""
    return [
        make_placement(
            "P001", "T001", "SUB1", "A", "ALL",
            ActivityType.LECTURE, [slot("MON", 1)], "R010",
        ),
        make_placement(
            "P002", "T002", "SUB2", "A", "ALL",
            ActivityType.LECTURE, [slot("MON", 2), slot("MON", 3)], "R011",
        ),
        make_placement(
            "P003", "T001", "SUB3", "A", "G1",
            ActivityType.PRACTICAL, [slot("TUE", 2), slot("TUE", 3)], "R001",
        ),
        make_placement(
            "P004", "T002", "SUB3", "A", "G2",
            ActivityType.PRACTICAL, [slot("TUE", 2), slot("TUE", 3)], "R002",
        ),
    ]


@pytest.fixture
def conflicted_placements() -> list:
    """Deliberate violations of C1, C2, C3 and C6."""
    return [
        # C1 teacher double-booking (T001 twice at MON P1)
        make_placement("P001", "T001", "SUB1", "A", "ALL",
                       ActivityType.LECTURE, [slot("MON", 1)], "R010"),
        make_placement("P002", "T001", "SUB2", "B", "ALL",
                       ActivityType.LECTURE, [slot("MON", 1)], "R011"),
        # C2 room double-booking (R010 at TUE P2, different sections)
        make_placement("P003", "T002", "SUB2", "A", "ALL",
                       ActivityType.LECTURE, [slot("TUE", 2)], "R010"),
        make_placement("P004", "T003", "SUB1", "B", "ALL",
                       ActivityType.LECTURE, [slot("TUE", 2)], "R010"),
        # C3 section conflict: section A lecture overlapping a G1 practical
        make_placement("P005", "T002", "SUB2", "A", "ALL",
                       ActivityType.LECTURE, [slot("WED", 2)], "R011"),
        make_placement("P006", "T003", "SUB3", "A", "G1",
                       ActivityType.PRACTICAL, [slot("WED", 2), slot("WED", 3)], "R001"),
        # C6 parallel misalignment: G1 and G2 at different times
        make_placement("P007", "T001", "SUB3", "B", "G1",
                       ActivityType.PRACTICAL, [slot("THU", 2), slot("THU", 3)], "R001"),
        make_placement("P008", "T002", "SUB3", "B", "G2",
                       ActivityType.PRACTICAL, [slot("FRI", 5), slot("FRI", 6)], "R002"),
    ]


# ====================================================================
# Time grid
# ====================================================================

class TestTimeGrid:

    def test_column_order_keeps_lunch_between_4_and_5(self):
        assert GRID_COLUMNS == ("1", "2", "3", "4", COL_LUNCH, "5", "6", "7")

    def test_period_times(self):
        assert PERIOD_TIMES[1] == ("9:00", "10:00")
        assert PERIOD_TIMES[4] == ("12:00", "1:00")
        assert PERIOD_TIMES[5] == ("2:00", "3:00")
        assert PERIOD_TIMES[7] == ("4:00", "5:00")

    def test_lunch_label(self):
        assert lunch_label() == "RECESS 1:00-2:00"

    def test_block_positions(self):
        assert slots_are_consecutive([slot("MON", 1), slot("MON", 2)])
        assert slots_are_consecutive([slot("MON", 6), slot("MON", 7)])
        # 4 → 5 crosses the lunch break: never a valid block
        assert not slots_are_consecutive([slot("MON", 4), slot("MON", 5)])
        # different days are never consecutive
        assert not slots_are_consecutive([slot("MON", 2), slot("TUE", 3)])


# ====================================================================
# Conflict audit
# ====================================================================

class TestConflictAudit:

    def test_clean_timetable_passes(self, clean_placements):
        report = audit_placements(clean_placements)
        assert report.is_clean
        assert report.conflicts == []

    def test_teacher_double_booking(self, conflicted_placements):
        report = audit_placements(conflicted_placements)
        teacher_conflicts = [c for c in report.conflicts
                             if c.conflict_type == "TEACHER"]
        assert len(teacher_conflicts) == 1
        assert teacher_conflicts[0].entity_id == "T001"
        assert teacher_conflicts[0].day == "MON"
        assert teacher_conflicts[0].period == 1

    def test_room_double_booking(self, conflicted_placements):
        report = audit_placements(conflicted_placements)
        room_conflicts = [c for c in report.conflicts
                          if c.conflict_type == "ROOM"]
        assert len(room_conflicts) == 1
        assert room_conflicts[0].entity_id == "R010"
        assert room_conflicts[0].day == "TUE"

    def test_section_group_conflict(self, conflicted_placements):
        report = audit_placements(conflicted_placements)
        section_conflicts = [c for c in report.conflicts
                             if c.conflict_type == "SECTION"]
        assert len(section_conflicts) == 1
        assert section_conflicts[0].entity_id == "A"

    def test_parallel_misalignment(self, conflicted_placements):
        report = audit_placements(conflicted_placements)
        parallel = [c for c in report.conflicts
                    if c.conflict_type == "PARALLEL"]
        assert len(parallel) == 1
        assert "G1" in parallel[0].description
        assert "G2" in parallel[0].description

    def test_parallel_pair_with_different_teachers_is_valid(
        self, clean_placements,
    ):
        """Same subject, G1+G2 at the same slots in different rooms with
        different teachers is the *correct* parallel arrangement."""
        report = audit_placements(clean_placements)
        assert report.is_clean

    def test_conflicted_slot_keys(self, conflicted_placements):
        report = audit_placements(conflicted_placements)
        keys = {(pid, day, period)
                for pid, day, period in __import__(
                    "app.export.conflicts", fromlist=["x"]
                ).conflicted_slot_keys(report)}
        assert ("P001", "MON", 1) in keys
        assert ("P002", "MON", 1) in keys

    def test_unscheduled_passthrough(self, clean_placements):
        unscheduled = [{"assignment_id": "A9", "reasons": ["no room"]}]
        report = audit_placements(clean_placements, unscheduled)
        assert report.unscheduled == unscheduled


# ====================================================================
# View builders
# ====================================================================

class TestSectionView:

    def test_parallel_pair_combined_entry(self, clean_placements, ctx):
        report = audit_placements(clean_placements)
        views = build_views(clean_placements, ctx, report, VIEW_SECTION)
        assert len(views) == 1
        view = views[0]
        assert view.entity_id == "CSE-3-A"

        cell = view.cell(1, 1)  # Tuesday, P2
        assert not cell.is_empty
        assert len(cell.blocks) == 1
        assert cell.blocks[0].span == 2
        assert cell.blocks[0].is_start
        # One combined entry: subject + G1 line + G2 line
        assert len(cell.entries) == 1
        lines = cell.entries[0].lines
        assert len(lines) == 3
        assert "3.7 Data Structures Lab" in lines[0]
        assert lines[1].startswith("G1 · RJS")
        assert lines[2].startswith("G2 · SRP")
        assert "(Practical)" in lines[1]

    def test_block_continuation_cells_marked(self, clean_placements, ctx):
        report = audit_placements(clean_placements)
        view = build_views(clean_placements, ctx, report, VIEW_SECTION)[0]
        start = view.cell(1, 1)
        cont = view.cell(1, 2)
        assert start.blocks[0].is_start
        assert not cont.blocks[0].is_start
        assert (cont.blocks[0].placement_id
                == start.blocks[0].placement_id)

    def test_weekly_periods(self, clean_placements, ctx):
        report = audit_placements(clean_placements)
        view = build_views(clean_placements, ctx, report, VIEW_SECTION)[0]
        assert view.weekly_periods == 1 + 2 + 2 + 2


class TestGroupView:

    def test_g1_view_shows_own_group_and_all(self, clean_placements, ctx):
        report = audit_placements(clean_placements)
        views = build_views(clean_placements, ctx, report, VIEW_GROUP)
        g1 = next(v for v in views if v.entity_id == "CSE-3-A-G1")

        # Monday P1 (ALL lecture) + Monday P2-P3 lecture block + Tuesday block
        assert not g1.cell(0, 0).is_empty       # lecture
        assert not g1.cell(1, 1).is_empty       # practical
        assert not g1.cell(0, 1).is_empty       # lecture block on Monday
        assert g1.cell(1, 1).blocks[0].span == 2
        assert g1.weekly_periods == 1 + 2 + 2

    def test_parallel_counterpart_note(self, clean_placements, ctx):
        report = audit_placements(clean_placements)
        views = build_views(clean_placements, ctx, report, VIEW_GROUP)
        g2 = next(v for v in views if v.entity_id == "CSE-3-A-G2")
        cell = g2.cell(1, 1)
        notes = [e for e in cell.entries if e.is_note]
        assert len(notes) == 1
        assert "Parallel G1" in notes[0].lines[0]
        assert "CC1" in notes[0].lines[0]

    def test_no_group_view_without_groupings(self, ctx):
        solo = make_placement("P001", "T001", "SUB1", "A", "ALL",
                              ActivityType.LECTURE, [slot("MON", 1)], "R010")
        report = audit_placements([solo])
        views = build_views([solo], ctx, report, VIEW_GROUP)
        assert views == []


class TestTeacherAndRoomViews:

    def test_teacher_view_aggregates_sections(self, clean_placements, ctx):
        report = audit_placements(clean_placements)
        views = build_views(clean_placements, ctx, report, VIEW_TEACHER)
        by_id = {v.entity_id: v for v in views}
        assert set(by_id) == {"T001", "T002"}

        t1 = by_id["T001"]
        assert "RJS" in t1.title
        # MON P1 lecture (sec A) + TUE block (G1 practical)
        assert not t1.cell(0, 0).is_empty
        assert not t1.cell(1, 1).is_empty
        assert t1.weekly_periods == 3
        # Cell shows the section and group
        tue_cell = t1.cell(1, 1)
        assert any("CSE-3-A · G1" in e.lines[1] for e in tue_cell.entries)

    def test_room_view_lists_used_rooms(self, clean_placements, ctx):
        report = audit_placements(clean_placements)
        views = build_views(clean_placements, ctx, report, VIEW_ROOM)
        ids = {v.entity_id for v in views}
        assert ids == {"R001", "R002", "R010", "R011"}
        # R010 hosts only the Monday lecture
        l5 = next(v for v in views if v.entity_id == "R010")
        assert l5.weekly_periods == 1


class TestClassView:

    def test_sections_stacked_in_class_cell(self, ctx):
        items = [
            make_placement("P001", "T001", "SUB1", "A", "ALL",
                           ActivityType.LECTURE, [slot("MON", 1)], "R010"),
            make_placement("P002", "T002", "SUB1", "B", "ALL",
                           ActivityType.LECTURE, [slot("MON", 1)], "R010"),
        ]
        report = audit_placements(items)
        views = build_views(items, ctx, report, VIEW_CLASS)
        assert len(views) == 1
        view = views[0]
        assert view.entity_id == "CSE-3"
        cell = view.cell(0, 0)
        assert len(cell.entries) == 2
        prefixes = sorted(e.lines[0][:4] for e in cell.entries)
        assert prefixes == ["A · ", "B · "]

    def test_class_view_overlapping_blocks_repeat_not_merge(self, ctx):
        """Two sections running practicals across the same periods: the
        blocks overlap each other, so neither may merge and BOTH stay
        visible in every covered column (no blank periods)."""
        items = [
            make_placement("PA1", "T001", "SUB3", "A", "G1",
                           ActivityType.PRACTICAL,
                           [slot("MON", 2), slot("MON", 3)], "R001"),
            make_placement("PA2", "T002", "SUB3", "A", "G2",
                           ActivityType.PRACTICAL,
                           [slot("MON", 2), slot("MON", 3)], "R002"),
            make_placement("PB1", "T003", "SUB3", "B", "G1",
                           ActivityType.PRACTICAL,
                           [slot("MON", 2), slot("MON", 3)], "R001"),
        ]
        report = audit_placements(items)
        view = build_views(items, ctx, report, VIEW_CLASS)[0]

        from app.export.pdf_exporter import _plan_merges
        assert _plan_merges(view) == {}   # overlap -> nothing merges

        for col in (1, 2):  # P2 and P3 both show all sections
            cell = view.cell(0, col)
            texts = [" ".join(e.lines) for e in cell.entries]
            assert any("A · " in t for t in texts), (col, texts)
            assert any("B · " in t for t in texts), (col, texts)

    def test_class_view_shows_parallel_split(self, clean_placements, ctx):
        report = audit_placements(clean_placements)
        views = build_views(clean_placements, ctx, report, VIEW_CLASS)
        cell = views[0].cell(1, 1)
        lines = cell.entries[0].lines
        assert lines[0].startswith("A · ")
        assert any(l.startswith("G1 ·") for l in lines)
        assert any(l.startswith("G2 ·") for l in lines)


# ====================================================================
# CSV exporters
# ====================================================================

class TestCsvExport:

    def test_section_csv_roundtrip(self, clean_placements, ctx, tmp_path):
        exporter = TimetableExporter(clean_placements, ctx)
        path = exporter.export_view("section", "csv", out_dir=tmp_path)

        with open(path, encoding="utf-8-sig", newline="") as f:
            rows = list(csv.reader(f))

        # Header row: Day + 7 period columns + lunch column
        header = next(r for r in rows if r and r[0] == "Day")
        assert len(header) == 9
        assert header[1].startswith("P1")
        assert header[5].startswith("RECESS") or header[5].startswith("LUNCH")

        monday = next(r for r in rows if r and r[0] == "Monday")
        assert "3.2 Data Structures" in monday[1]
        assert monday[5] == "RECESS 1:00-2:00"

        tuesday = next(r for r in rows if r and r[0] == "Tuesday")
        # Block content repeats in each covered period (P2 and P3)
        assert "G1 · RJS" in tuesday[2]
        assert "G2 · SRP" in tuesday[2]
        assert "G1 · RJS" in tuesday[3]

    def test_flat_csv_one_row_per_slot(self, clean_placements, ctx, tmp_path):
        exporter = TimetableExporter(clean_placements, ctx)
        path = exporter.export_aux("flat", out_dir=tmp_path)

        with open(path, encoding="utf-8-sig", newline="") as f:
            rows = list(csv.DictReader(f))

        total_slots = sum(len(p.slots) for p in clean_placements)
        assert len(rows) == total_slots
        assert all(r["conflicts"] == "" for r in rows)

        block_rows = [r for r in rows if r["placement_id"] == "P003"]
        assert len(block_rows) == 2
        assert block_rows[0]["block_slot"] == "1/2"
        assert block_rows[1]["block_slot"] == "2/2"
        assert block_rows[0]["room_name"] == "CC1"

    def test_conflicts_csv_clean(self, clean_placements, ctx, tmp_path):
        exporter = TimetableExporter(clean_placements, ctx)
        path = exporter.export_aux("conflicts", out_dir=tmp_path)
        with open(path, encoding="utf-8-sig", newline="") as f:
            rows = list(csv.reader(f))
        assert rows[0][:2] == ["type", "entity_id"]
        assert rows[1][0] == "NONE"

    def test_conflicts_csv_dirty(self, conflicted_placements, ctx, tmp_path):
        exporter = TimetableExporter(conflicted_placements, ctx)
        path = exporter.export_aux("conflicts", out_dir=tmp_path)
        with open(path, encoding="utf-8-sig", newline="") as f:
            rows = list(csv.DictReader(f))
        types = {r["type"] for r in rows}
        assert types == {"TEACHER", "ROOM", "SECTION", "PARALLEL"}

    def test_unscheduled_csv(self, clean_placements, ctx, tmp_path):
        exporter = TimetableExporter(
            clean_placements, ctx,
            unscheduled=[{"assignment_id": "A9", "subject_id": "SUB1",
                          "section": "A", "group": "ALL",
                          "activity_type": "LECTURE",
                          "reasons": ["no free slot"]}],
        )
        path = exporter.export_aux("unscheduled", out_dir=tmp_path)
        with open(path, encoding="utf-8-sig", newline="") as f:
            rows = list(csv.DictReader(f))
        assert len(rows) == 1
        assert rows[0]["assignment_id"] == "A9"
        assert "no free slot" in rows[0]["reasons"]

    def test_conflict_flag_in_grid_csv(
        self, conflicted_placements, ctx, tmp_path,
    ):
        exporter = TimetableExporter(conflicted_placements, ctx)
        path = exporter.export_view("section", "csv", out_dir=tmp_path)
        text = path.read_text(encoding="utf-8-sig")
        assert "⚠ CONFLICT" in text


# ====================================================================
# XLSX exporters
# ====================================================================

class TestXlsxExport:

    def test_structure_and_merges(self, clean_placements, ctx, tmp_path):
        import openpyxl

        exporter = TimetableExporter(clean_placements, ctx)
        path = exporter.export_view("section", "xlsx", out_dir=tmp_path)
        wb = openpyxl.load_workbook(path)

        assert wb.sheetnames == ["CSE-3-A", "Conflicts"]
        ws = wb["CSE-3-A"]
        merges = {str(m) for m in ws.merged_cells.ranges}
        # 2-slot lecture block at Monday P2-P3 and practical at Tue P2-P3
        assert "C5:D5" in merges    # Monday P2:P3
        assert "C6:D6" in merges    # Tuesday P2:P3
        assert "F5:F9" in merges    # lunch column

        # Header labels carry readable times
        assert ws.cell(row=4, column=2).value.startswith("P1")

        # Parallel pair renders both groups in one cell
        tuesday_text = ws.cell(row=6, column=3).value
        assert "G1 · RJS" in tuesday_text
        assert "G2 · SRP" in tuesday_text

    def test_conflicts_sheet_clean(self, clean_placements, ctx, tmp_path):
        import openpyxl

        exporter = TimetableExporter(clean_placements, ctx)
        path = exporter.export_view("teacher", "xlsx", out_dir=tmp_path)
        wb = openpyxl.load_workbook(path)
        ws = wb["Conflicts"]
        assert "No conflicts detected" in ws["A1"].value

    def test_conflict_cells_flagged(self, conflicted_placements, ctx, tmp_path):
        import openpyxl

        exporter = TimetableExporter(conflicted_placements, ctx)
        path = exporter.export_view("class", "xlsx", out_dir=tmp_path)
        wb = openpyxl.load_workbook(path)

        assert "Conflicts" in wb.sheetnames
        ws = wb["Conflicts"]
        conflict_rows = [
            r for r in range(4, ws.max_row + 1)
            if ws.cell(row=r, column=1).value in
            ("TEACHER", "ROOM", "SECTION", "PARALLEL")
        ]
        assert len(conflict_rows) == 4

        # The offending class-view cell is marked with ⚠
        class_ws = wb["CSE-3"]
        monday_p1 = class_ws.cell(row=5, column=2)
        assert monday_p1.value and "⚠" in monday_p1.value

    def test_sheet_names_sanitized(self, ctx):
        from app.export.xlsx_exporter import _sheet_name
        from app.export.views import GridView

        view = GridView(VIEW_SECTION, "weird/name*:?", "t", "s")
        name = _sheet_name(view, set())
        assert not (set(name) & set('/\\*:?[]'))


# ====================================================================
# PDF exporters
# ====================================================================

class TestPdfExport:

    @pytest.mark.parametrize("view_kind", list(VIEW_KINDS))
    def test_valid_pdf_per_view(
        self, clean_placements, ctx, tmp_path, view_kind,
    ):
        exporter = TimetableExporter(clean_placements, ctx)
        path = exporter.export_view(view_kind, "pdf", out_dir=tmp_path)

        raw = path.read_bytes()
        assert raw[:5] == b"%PDF-"
        assert len(raw) > 2000

    def test_page_count_matches_entities(self, clean_placements, ctx, tmp_path):
        import pymupdf

        exporter = TimetableExporter(clean_placements, ctx)
        path = exporter.export_view("teacher", "pdf", out_dir=tmp_path)
        doc = pymupdf.open(str(path))
        # 1 audit page + 2 teachers
        assert len(doc) == 3

    def test_conflicts_listed_on_audit_page(
        self, conflicted_placements, ctx, tmp_path,
    ):
        import pymupdf

        exporter = TimetableExporter(conflicted_placements, ctx)
        path = exporter.export_view("section", "pdf", out_dir=tmp_path)
        doc = pymupdf.open(str(path))
        text = doc[0].get_text()
        assert "CONFLICT AUDIT" in text
        assert "TEACHER" in text and "PARALLEL" in text


# ====================================================================
# Export service
# ====================================================================

class TestExportService:

    def test_export_all_bundle(self, clean_placements, ctx, tmp_path):
        exporter = TimetableExporter(clean_placements, ctx)
        result = exporter.export_all(out_dir=tmp_path)

        # 5 views × 3 formats + flat + conflicts + unscheduled
        assert len(result.files) == 18
        names = {f.name for f in result.files}
        assert "Timetable_Class.xlsx" in names
        assert "Timetable_Room.pdf" in names
        assert "Timetable_Flat.csv" in names
        assert "Conflicts.csv" in names
        assert result.is_clean

    def test_entity_filter(self, clean_placements, ctx, tmp_path):
        exporter = TimetableExporter(clean_placements, ctx)
        path = exporter.export_view(
            "teacher", "csv", entity="T001", out_dir=tmp_path,
        )
        assert path.name == "Timetable_Teacher_T001.csv"
        text = path.read_text(encoding="utf-8-sig")
        assert "RJS (T001)" in text
        assert "SRP" not in text

    def test_entity_filter_unknown_raises(self, clean_placements, ctx, tmp_path):
        exporter = TimetableExporter(clean_placements, ctx)
        from app.export.service import ExportError

        with pytest.raises(ExportError):
            exporter.export_view("teacher", "csv", entity="T999",
                                 out_dir=tmp_path)

    def test_default_filename(self, clean_placements, ctx):
        exporter = TimetableExporter(clean_placements, ctx)
        assert exporter.default_filename("section", "xlsx") == \
            "Timetable_Section.xlsx"
        assert exporter.default_filename("section", "pdf", "CSE-3-A") == \
            "Timetable_Section_CSE-3-A.pdf"


# ====================================================================
# JSON rebuild (server restart fallback)
# ====================================================================

class TestRebuild:

    def test_placements_from_result_data(self, clean_placements, ctx):
        result_data = {
            "placements": [
                {
                    "placement_id": p.placement_id,
                    "assignment_id": p.assignment.assignment_id,
                    "teacher_id": p.assignment.teacher_id,
                    "subject_id": p.assignment.subject_id,
                    "section": p.assignment.section,
                    "group": p.assignment.group,
                    "activity_type": p.assignment.activity_type.value,
                    "room_id": p.room_id,
                    "room_name": p.room_id,
                    "slots": [
                        {"day": s.day.name, "period": s.period}
                        for s in p.slots
                    ],
                }
                for p in clean_placements
            ],
        }

        rebuilt = placements_from_result_data(result_data, ctx)
        assert len(rebuilt) == len(clean_placements)
        for orig, new in zip(clean_placements, rebuilt):
            assert new.assignment.teacher_id == orig.assignment.teacher_id
            assert new.assignment.section == orig.assignment.section
            assert new.assignment.group == orig.assignment.group
            assert new.room_id == orig.room_id
            assert len(new.slots) == len(orig.slots)

        # Rebuilt placements export cleanly (same audit result)
        exporter = TimetableExporter(rebuilt, ctx)
        assert exporter.report.is_clean

    def test_conflicts_survive_rebuild(self, conflicted_placements, ctx):
        result_data = {
            "placements": [
                {
                    "placement_id": p.placement_id,
                    "teacher_id": p.assignment.teacher_id,
                    "subject_id": p.assignment.subject_id,
                    "section": p.assignment.section,
                    "group": p.assignment.group,
                    "activity_type": p.assignment.activity_type.value,
                    "room_id": p.room_id,
                    "slots": [
                        {"day": s.day.name, "period": s.period}
                        for s in p.slots
                    ],
                }
                for p in conflicted_placements
            ],
        }
        rebuilt = placements_from_result_data(result_data, ctx)
        report = audit_placements(rebuilt)
        assert not report.is_clean
        assert set(report.count_by_type()) == {
            "TEACHER", "ROOM", "SECTION", "PARALLEL",
        }
